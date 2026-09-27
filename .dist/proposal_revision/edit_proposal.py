from docx import Document
from docx.shared import RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from pathlib import Path

ROOT=Path('E:/Code/productSelectionAgent')
doc=Document(ROOT/'Shopping_Copilot_Proposal_EN.docx')
def find(prefix):
    return next(p for p in doc.paragraphs if p.text.startswith(prefix))
def replace(prefix,text):
    p=find(prefix); p.text=text; return p
def before(anchor,text,style='Normal'):
    p=anchor.insert_paragraph_before(text,style)
    return p
def add_block(anchor,blocks):
    for style,text in blocks: before(anchor,text,style)

replace('Pilot targets are a reduction', 'Pilot targets are a reduction in median time to a usable shortlist from an assumed 10 to 6 minutes, a 40% reduction in repetitive selection workload, and 98% evidence accuracy for sampled objective claims. Conversion growth from 3.0% to 3.3% remains a hypothesis. The prototype supports a synthetic-catalog demonstration and a separate real-search adapter. Local search data, index and model files are present; their presence alone does not establish readiness or data quality. Confirmation records a session decision, not an order or payment. [S3–S7]')
replace('Existing scope includes', 'Existing scope includes session requirements, catalog retrieval and reranking, evidence displays, product questions, saving, comparison, explicit confirmation and audit export. Search currently explores one product category at a time; deferred category details can be remembered for later. Pilot work adds retailer data acceptance, protected deployment, human assistance and business-event/order attribution. Accounts and cross-session personalization remain future work. Live web-wide price comparison, stock reservation, guaranteed fit, payment, refunds and autonomous ordering are outside scope. [S2–S8]')
add_block(find('Scope and subsequent work'),[
('Heading 2','How the functions work together'),
('Normal','A message first resolves the shopper’s intent and any reference to the visible products. The workflow applies explicit requirement changes, then chooses a bounded next action: clarify, search, answer, compare, save or confirm. “How much is #2?” reads the current item without searching again; “Blue instead” changes the active requirement and refreshes eligible results. Undo restores the prior state and relevant view. This continuity is the mechanism intended to reduce repeated explanation and page checking.'),
('Normal','For example, “I need a blue dress. I prefer cotton” records dress and blue as requirements and cotton as a preference. Retrieval finds candidates; catalog checks enforce requirements, while supported cotton matches can move ahead without excluding all alternatives. “Compare #1 and #2” binds those numbers to stable product IDs before comparison. The shopper can revise the preference, compare again and explicitly finalize a saved shortlist. Unknown price or stock stays unknown throughout. [S3–S6]')])

# Replace the technical section in place, keeping the source document's section hierarchy.
start=find('Existing architecture and data flow'); end=find('08  Risk management')
cur=start._p
while cur is not end._p:
    nxt=cur.getnext(); cur.getparent().remove(cur); cur=nxt
blocks=[
('Heading 2','Architecture and reasoning loop'),
('Normal','Shopping Copilot uses a stateful Python orchestrator with specialist modules for intent, memory, retrieval, ranking and comparison. Its planning pattern is an adaptive policy loop: interpret the message and current context, validate a state update, choose a permitted action, execute the relevant tool, inspect its result, and return evidence with the next available choice. Candidate evidence can justify a clarification; a failure can produce a recoverable response. The loop advances through shopper turns rather than running an unrestricted autonomous task. Generative assistance is optional and off by default. [S3, S6]'),
('Normal','End to end: shopper message → intent and reference resolution → versioned requirements → policy decision → search or a read-only answer or comparison → evidence and visible options → shopper feedback → updated state. Saving and finalization are explicit branches. Deterministic controls retain responsibility for state mutation and action boundaries even when a model assists interpretation.'),
('Heading 2','Intent recognition and requirement updates'),
('Normal','The intent router and control handlers distinguish browsing, directed search, comparison, product questions and commands such as Undo or Start over. Catalog-aware extraction separates hard requirements, weighted soft preferences and exclusions. Separately phrased questions and edits can share a turn; product references bind to the pre-update display. Ambiguous categories or item references request clarification rather than silently selecting an interpretation. This produces an actionable request while preserving what the shopper actually committed to. The language coverage is bounded, not unrestricted understanding. [S3, S6]'),
('Heading 2','Explicit state and memory'),
('Normal','Session memory stores the active category, hard conditions, weighted preferences, exclusions, shown and rejected product IDs, pending questions and deferred category details. Source turns and evidence make requirement changes inspectable. Snapshots and history support Undo and Redo, including the relevant visible results; saved products use stable IDs rather than changing screen ranks. A session ID, turn and state version travel with retrieval results so stale or cross-session results can be rejected. The effect is continuity through corrections without replaying the whole conversation. [S3, S6]'),
('Normal','The service keeps session data in memory, with defaults of 128 sessions and a 3,600-second idle lifetime. These are storage limits, not measured concurrent-user capacity. The proposal claims session memory only. Durable cross-device memory, account ownership and shared storage require additional implementation.'),
('Heading 2','Catalog retrieval and contextual ranking'),
('Normal','The search adapter constructs an English query from the structured requirements and calls search_products(query, top_k). The shared search tool retrieves lexical candidates with BM25, then scores query–product pairs using a local sequence-classification model. get_product_details(product_ids) supplies listing fields for those IDs. The adapter validates the returned IDs and finite scores, checks supported hard conditions and exclusions, and favors unseen items when more results are requested. The interface shows at most ten products. [S4]'),
('Normal','Catalog-supported soft preferences for attributes such as color, material or fit can adjust the final order without becoming hard exclusions. The underlying search scores remain relevance scores, not probabilities, quality guarantees or purchase confidence. When evidence is missing, the system does not invent a match. The current adapter leaves price and ratings unknown, so a strict budget cannot be verified against those records. An empty pool prompts an explicit revision or recovery choice; constraints are not silently relaxed. The intended effect is a smaller, relevant and explainable set of options. [S3, S4]'),
('Heading 2','Evidence based clarification and product questions'),
('Normal','The adaptive policy uses distinctions in retrieved candidates to decide whether an optional question could meaningfully narrow the choice. Its attention limits reduce repeated questioning; explicit browsing or a request to see results can skip optional refinement. A missing searchable category or unverifiable strict budget requires a different clarification. Questions about a displayed item use its catalog record and preserve the current page and open choice. Unsupported attributes are reported as unknown. This keeps guidance useful while avoiding unnecessary searches and forced questionnaires. [S3]'),
('Heading 2','Description and comparison'),
('Normal','The description adapter passes up to three resolved products and current requirements to DescriptionComparison. With a configured model provider, the pipeline extracts category-relevant attributes, creates an objective comparison, and then adds personalized suitability reasons and cautions. Personalization is skipped when no requirements or preferences exist. Source-quote checks reject unsupported quotations; inferred values remain labeled. Objective facts are not rewritten to fit a preference, and comparison preserves the requested product order. [S5]'),
('Normal','Without a provider, offline_preview exposes direct catalog facts and unknowns, without generated personalized advice. A failed model stage yields partial output; an explicit retry can resume the comparison. The session cache reuses completed comparisons when product IDs and order, requirements, supplied facts and model identity are unchanged. Changed inputs invalidate the cache. This avoids unnecessary repeat calls while keeping updated comparisons aligned with the shopper’s current needs.'),
('Heading 2','Typed tool boundaries and module integration'),
('Normal','RetrievalResultV2 and RankingResultV2 define versioned contracts for session context, candidates and ranking. Validators check identity, candidate limits, continuous ranks and matching state context. The selection handoff uses show-me-your-agent.selection.v1; comparison returns description-comparison.v1 with product profiles, comparison values, evidence references, status and usage. Empty, duplicate or oversized selections are handled explicitly. These contracts make handoffs inspectable and reduce stale-data and malformed-output errors. [S3–S5]'),
('Heading 2','Platform and tooling usage'),
('Normal','The implementation uses Python modules and explicit orchestration rather than a framework-managed team of autonomous LLM agents. BM25, PyTorch and Transformers serve local search; JSON interfaces connect the service and shopper UI. A configurable provider supplies optional local or cloud model assistance. Adapters connect the team’s search and description modules without requiring their internals to share conversation state. This separation supports independent module testing and a consistent state authority. [S3–S6]'),
('Heading 2','Prerequisites and scaling'),
('Normal','The retailer must supply authorized records with stable IDs, source fields, price currency, timestamps and variant or stock evidence. The local search table, index and model files exist, but the review runtime lacks several search dependencies; live retrieval was not rerun in this review. A pilot needs a configured runtime and real-query acceptance, not just a passing health check. The current catalog mapping also needs real price and inventory integration. [S4, S7]'),
('Normal','Before public traffic, add session ownership, rate limits, production API routing, shared state, per-session concurrency controls, timeouts and monitored fallback. The prototype HTTP service and serialization are not proof of production throughput. Measure latency and token usage separately for search, model stages and cache hits; load-test at twice the retailer’s measured peak. Additional categories need schema and evidence tests; additional retailers need tenant isolation. These are deployment requirements, not existing capacity claims. [S8]')]
add_block(end,blocks)
before(end,'').add_run().add_break(__import__('docx').enum.text.WD_BREAK.PAGE)

add_block(find('Experiment and attribution'),[
('Heading 2','Observability and technical evaluation'),
('Normal','Per-turn receipts expose the chosen action and reason, requirement changes, result IDs, evidence and model usage. Audit export links execution records with digest checks; it helps inspect sequence consistency but is not a tamper-proof external ledger. Comparison also reports status and latency. Reviewers can follow why a clarification, retrieval or confirmation occurred through structured decisions rather than private model reasoning. Production retention, access controls and operational dashboards remain pilot work. [S3, S5]'),
('Normal','Golden-path evaluation follows a need through retrieval, correction, Undo, saving, comparison and explicit confirmation. Failure cases cover empty results, unpriced strict budgets, invalid product references, stale state, search failure, partial comparison and changes after confirmation. Before release, extend adversarial cases with instructions embedded in product text, fabricated evidence, cross-session requests and repeated or oversized input. Check factual support, requirement preservation, recovery and absence of unintended confirmation. Record elapsed task time and manual steps against the same baseline task; automated checks alone cannot establish business savings. [S9]')])

add_block(find('Human assistance workflow'),[
('Heading 2','Autonomy and approval checkpoints'),
('Normal','The agent can retrieve, rank, display evidence and propose useful questions within the session. The shopper controls preference changes, saved options and finalization; a generic acknowledgement or conditional instruction must not be treated as permission to finalize. Edits can invalidate a previous confirmation. No purchase, payment or stock reservation tool is available. The demo must show this real confirmation boundary. Staff escalation for exceptional or commercially binding decisions belongs to the pilot workflow below and is not an implemented approval service. [S3]'),
('Heading 2','Safety and security boundaries'),
('Normal','The description prompt explicitly treats product text as untrusted data, and the pipeline checks supporting quotes against supplied product fields. Typed inputs, bounded product selection and limited tool access constrain what model output can affect. These controls reduce exposure but do not prove prompt-injection immunity or correct interpretation of every quote. Before deployment, test malicious catalog instructions and enforce least privilege through backend credentials, session ownership, authorized catalog access and protected logs. Cloud model use must disclose that messages and selected catalog excerpts may leave the local service; avoid sending unnecessary personal information. [S3, S5, S6]')])

replace('Use a 5–6 minute demo', 'Use a 5–6 minute demo with one business story: turning a changing shopping need into an evidence-backed shortlist. Show the synthetic Demo catalog label and start a new session. Keep a visible timer and show the same task in the baseline journey when comparing manual steps. Use the configured interface’s equivalent controls or chat commands. A second, explicitly labeled model-enabled segment is needed to demonstrate generated personalized advice. [S2, S3]')
replace('A timed demonstration can show', 'A timed demonstration shows the observed task duration, tool sequence and approval point; it does not establish a population-wide 40% reduction. Show requirement state before and after a correction, the resolved comparison IDs, an unknown attribute, and the finalization record. Demonstrate one invalid reference or unpriced strict-budget recovery without inventing data or silently relaxing requirements. Keep staff handoff labeled as planned. With the model off, describe comparison as an offline facts preview, not validated generative reasoning. [S3–S6]')

# Remove obsolete external-workspace references while keeping an auditable source register.
replacements={
'References identify the current':'References identify the local project modules reviewed for this proposal. S1 maps the judging requirements; S2–S8 support architecture and capability statements; S9 states the verification boundary. Business baselines, budgets, growth figures and service commitments remain planning assumptions or targets.',
'S1  ':'S1  Supplied judging slides and business proposal guidelines\nThe judging slide covers goal and scope, architecture and reasoning, tools, human oversight, safety, observability and platform usage. The Consistent Message slide requires proposal claims to match demo evidence. These map to Sections 1–3, 7, 8, 5 and 10.',
'S2  ':'S2  agentic_workflow/SUBMISSION_EN.md and DESCRIPTION_INTEGRATION.md\nDocumented shopper journeys, limitations and integration behavior; checked against the runtime and adapters. Historical verification counts are not new measurements.',
'S3  ':'S3  agentic_workflow/mvp/server.py; shopping_agent/agent.py; shopping_agent/policy.py; mvp/audit.py\nOrchestration, clarification, product references, session limits, selection, confirmation and execution records. Paths after the semicolon are relative to agentic_workflow.',
'S4  ':'S4  search_tool/tool.py; agentic_workflow/shopping_agent/search_adapter.py; agentic_workflow/retrieval-and-reranking/techjam_agent/contracts_v2.py\nBM25, local reranking, details lookup, catalog filtering, soft-preference ordering and versioned contracts.',
'S5  ':'S5  description_module/pipeline.py and README.md; agentic_workflow/shopping_agent/description_adapter.py\nExtraction, objective and personalized comparison, quote checks, statuses, handoff and caching integration.',
'S6  ':'S6  agentic_workflow/intent-recognition/intent_router; conversation-state-memory/src/state_memory; shopping_agent/requirement_enhancer.py; shopping_agent/model_provider.py; __main__.py\nIntent, structured memory and optional model configuration. Paths after the semicolon are relative to agentic_workflow.',
'S7  ':'S7  agentic_workflow/preflight.py; search_resources.py; search_tool/artifacts and search_tool/models\nLocal resource inspection found the product table, retrieval index and saved model. Preflight in the review runtime reports missing bm25s, torch, transformers and pyarrow dependencies; this is not evidence that all user environments lack them.',
'S8  ':'S8  agentic_workflow/mvp/server.py and deployment requirements in this proposal\nThe local service is a prototype. Authentication, shared state, hardened hosting, human assistance and retailer order attribution require implementation and acceptance. Separate frontend build claims are outside this review.',
'S9  ':'S9  Verification for this revision\nDescription module: 8 automated tests passed with simulated responses. Workflow regression result is recorded below. Search dependency preflight was run; no live-model quality, retailer-data acceptance or business-outcome measurement was performed.',
'Verification boundary:':'Verification boundary: offline and simulated tests establish covered behavior only. They do not establish retailer-data accuracy, live-model quality, prompt-injection immunity, concurrent capacity or business returns. Synthetic demonstration data must remain visibly identified; live-search and cloud-model acceptance require a separately configured run.'}
for prefix,text in replacements.items(): replace(prefix,text)

for style_name in ('Title','Subtitle','Heading 1','Heading 2','Heading 3'):
    doc.styles[style_name].font.color.rgb=RGBColor(0,0,0)
for p in doc.paragraphs:
    if p.style.name.startswith('Heading') or p.style.name in ('Title','Subtitle'):
        for r in p.runs: r.font.color.rgb=RGBColor(0,0,0)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs: p.paragraph_format.widow_control=True

out=ROOT/'Shopping_Copilot_Proposal_EN_Completed.docx'
doc.save(out)
print(out)
